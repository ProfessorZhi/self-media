#!/usr/bin/env python3
"""Extract published job cutoffs, never candidate scores or candidate identity.

Usage: python scripts/build_interview_cutoffs_2026.py --workspace /path/to/inputs
Requires openpyxl and xlrd. Reads the three pinned local inputs below; no network.
The national source is a public mirror. The Audit Office attachment independently
checks a limited subset at its own publication batch. Original XLS/XLSX files are
not copied into the repository. JSON outputs contain job-level data only.
"""

import argparse
import collections
import csv
import hashlib
import json
import re
import warnings
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import xlrd

YEAR = 2026
DATASET = 'national-civil-service-initial-interview-cutoffs-2026'
SHARD_LIMIT = 840_000  # Strict UTF-8 byte ceiling below the requested 850 KB.
RESEARCHED_AT = '2026-10-16T19:28:18Z'
SOURCES = [
    {
        'id': 'national_2026_initial_interview_mirror',
        'filename': '2026_interview_initial_mirror.xlsx',
        'url': 'https://u3.huatu.com/uploads/soft/260115/f1_2.xlsx',
        'publication_url': 'https://m.huatu.com/2026/0115/2846698.html',
        'published_at': '2026-01-15T11:31:00+08:00',
        'publisher': '华图教育；页面标注来源国家公务员局',
        'source_level': 'third_party_mirror_of_official_published_list',
        'batch_id': 'initial_2026_01_15',
        'batch_label': '2026年度国考首批面试名单（2026-01-15镜像发布）',
        'bytes': 4_825_937,
        'sha256': '6fa1684a40a97c0094b42a0ab43d53742a2e2293a9ff12f37bf3b09af194becd',
        'role': 'national_published_job_cutoffs',
    },
    {
        'id': 'audit_office_2026_interview_official',
        'filename': '2026_interview_audit_official.xls',
        'url': 'https://www.audit.gov.cn/n8/n28/c10764531/part/10764506.xls',
        'publication_url': 'https://www.audit.gov.cn/n8/n28/c10764531/content.html',
        'published_at': '2026-01-26T18:20:00+08:00',
        'publisher': '审计署人事教育司',
        'source_level': 'recruiting_department_official_attachment',
        'batch_id': 'audit_department_interview_2026_01_26',
        'batch_label': '审计署部门面试公告附件（2026-01-26）',
        'batch_note': '独立保留部门公告日期；不预先假定与全国首批、调剂或递补名单同批。',
        'bytes': 80_384,
        'sha256': '6cead9a2abb2449a7400dddc3de5e34eb0868d70ce0572b2540b49a211fe5882',
        'role': 'independent_official_subset_comparison',
    },
    {
        'id': 'upstream_csv_2026',
        'filename': '2026_position_source.csv',
        'url': 'https://raw.githubusercontent.com/icy-bean/civil-exam/main/clean/national/2026_position.csv',
        'publisher': 'icy-bean/civil-exam',
        'source_level': 'third_party_cleaned_position_table',
        'published_at': None,
        'bytes': 17_494_070,
        'sha256': '2186edb16faeca9aa7ab237c270c60e0c8092e25b20726f88ece7735481301b6',
        'role': 'composite_key_join_for_position_attributes',
        'expected_positions': 20_714,
        'expected_planned_seats': 38_119,
    },
]

# Explicitly reviewed official workbook unit labels, not a fuzzy job-code join.
AUDIT_UNITS = {
    '署机关': '审计署',
    '京津冀特派办': '审计署京津冀特派员办事处',
    **{f'{city}特派办': f'审计署驻{city}特派员办事处' for city in [
        '太原', '哈尔滨', '上海', '武汉', '广州', '郑州', '济南', '西安',
        '昆明', '长沙', '深圳', '长春', '重庆',
    ]},
}

CSV_FIELDS = [
    ('source_csv_province', 'work_province'),
    ('source_csv_city', 'work_city'),
    ('source_system_tag', 'treatment_system_class'),
    ('source_csv_work_location', 'work_location'),
    ('source_csv_education', 'education_raw'),
    ('source_csv_degree', 'degree_raw'),
    ('source_csv_major_raw', 'major_raw'),
    ('source_csv_major_codes_raw', 'major_codes'),
    ('source_csv_exam_category', 'exam_category'),
    ('source_csv_org_nature', 'org_nature'),
    ('source_csv_org_level', 'org_level'),
    ('source_csv_position_attribute', 'position_attr'),
    ('source_csv_interview_ratio', 'interview_ratio'),
    ('source_csv_planned_seats', 'headcount'),
]
COLUMNS = [
    'id', 'year', 'batch_id', 'source_id', 'source_sheet',
    'source_first_row', 'source_last_row', 'source_list_row_count',
    '招录机关', '部门代码', '用人司局', '招考职位', '职位代码',
    '最低面试分数', 'published_cutoff_values', 'source_cutoff_field',
    'score_unit', 'source_number_formats', 'cutoff_unique_in_source',
    'position_csv_match', 'position_csv_record',
    *[name for name, _ in CSV_FIELDS],
    'official_audit_comparison',
]


def check(condition, message):
    if not condition:
        raise ValueError(message)


def encode(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False).encode('utf-8')


def digest(data):
    return hashlib.sha256(data).hexdigest()


def write_json(path, value, pretty=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    data = (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False).encode('utf-8') if pretty else encode(value)) + b'\n'
    check(len(data) < SHARD_LIMIT, f'Output exceeds byte limit: {path.name}')
    path.write_bytes(data)
    return {'path': path.name, 'bytes': len(data), 'sha256': digest(data)}


def text(value):
    return '' if value is None else str(value).strip()


def norm(value):
    return re.sub(r'\s+', '', text(value))


def code(value, width):
    s = text(value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        check(float(value).is_integer(), 'Non-integer code in source')
        s = str(int(value))
    check(re.fullmatch(r'\d{1,' + str(width) + r'}', s) is not None, 'Invalid code structure in source')
    return s.zfill(width)


def key(dept, position):
    return f'{YEAR}:{dept}:{position}'


def score(value):
    if value is None or text(value) == '':
        return None
    try:
        result = Decimal(text(value))
    except InvalidOperation as exc:
        raise ValueError('Non-numeric value in published cutoff column') from exc
    check(result.is_finite() and result >= 0, 'Invalid published cutoff value')
    return result


def score_number(value):
    return int(value) if value == value.to_integral_value() else float(value)


def verify_inputs(workspace):
    for spec in SOURCES:
        path = workspace / spec['filename']
        data = path.read_bytes()
        check(len(data) == spec['bytes'], f'{path.name}: unexpected byte length')
        check(digest(data) == spec['sha256'], f'{path.name}: input hash mismatch')


def read_positions(workspace):
    positions = {}
    departments = collections.defaultdict(set)
    with (workspace / SOURCES[2]['filename']).open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        for row_number, row in enumerate(reader, 2):
            check(row['year'] == str(YEAR), 'Unexpected CSV recruitment year')
            dept, position = code(row['dept_code'], 6), code(row['position_code'], 12)
            pid = key(dept, position)
            check(pid not in positions, f'Duplicate position CSV key: {pid}')
            positions[pid] = {'record': row_number, 'data': row}
            departments[row['dept_name']].add(dept)
    check(len(positions) == SOURCES[2]['expected_positions'], 'Unexpected position CSV record count')
    check(sum(int(v['data']['headcount']) for v in positions.values()) == SOURCES[2]['expected_planned_seats'], 'Unexpected CSV plan total')
    return positions, departments


def read_national(workspace):
    groups = {}
    metadata = {'sheets': [], 'source_list_rows': 0, 'rows_with_missing_cutoff': 0}
    expected = ['准考证号', '姓名', '招录机关', '部门代码', '用人司局', '招考职位', '职位代码', '最低面试分数']
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='Workbook contains no default style')
        workbook = openpyxl.load_workbook(workspace / SOURCES[0]['filename'], read_only=True, data_only=True)
    try:
        for sheet in workbook:
            sheet_meta = {'sheet': sheet.title, 'declared_max_row': sheet.max_row, 'declared_max_column': sheet.max_column, 'dimensions_reset': True}
            sheet.reset_dimensions()
            rows = sheet.iter_rows()
            headers = [text(c.value) for c in next(rows)]
            check(headers == expected, 'Unexpected national interview list header')
            sheet_meta['header_row'] = 1
            sheet_meta['data_rows'] = 0
            for row_number, cells in enumerate(rows, 2):
                # Identity columns 0 and 1 are never projected, logged or persisted.
                if len(cells) < 8 or not text(cells[6].value):
                    continue
                dept, position = code(cells[3].value, 6), code(cells[6].value, 12)
                pid = key(dept, position)
                cutoff = score(cells[7].value)
                safe = {
                    '招录机关': text(cells[2].value), '部门代码': dept,
                    '用人司局': text(cells[4].value), '招考职位': text(cells[5].value),
                    '职位代码': position,
                }
                if pid not in groups:
                    groups[pid] = {'id': pid, 'source_sheet': sheet.title, 'first_row': row_number, 'last_row': row_number, 'row_count': 0, 'missing_score_rows': 0, 'scores': set(), 'formats': set(), 'raw': safe, 'label_variants': set()}
                group = groups[pid]
                check(group['source_sheet'] == sheet.title, f'Position spans sheets: {pid}')
                if safe != group['raw']:
                    group['label_variants'].add((safe['招录机关'], safe['用人司局'], safe['招考职位']))
                group['last_row'] = row_number
                group['row_count'] += 1
                group['formats'].add(cells[7].number_format)
                if cutoff is None:
                    group['missing_score_rows'] += 1
                    metadata['rows_with_missing_cutoff'] += 1
                else:
                    group['scores'].add(cutoff)
                sheet_meta['data_rows'] += 1
                metadata['source_list_rows'] += 1
            metadata['sheets'].append(sheet_meta)
    finally:
        workbook.close()
    check(metadata['source_list_rows'] == 113_377, 'National source list rows do not match published mirror total')
    return groups, metadata


def merged_value(sheet, row, column, merges):
    for r0, r1, c0, c1 in merges:
        if r0 <= row < r1 and c0 <= column < c1:
            return sheet.cell_value(r0, c0), r0
    return sheet.cell_value(row, column), row


def read_audit(workspace, departments):
    workbook = xlrd.open_workbook(workspace / SOURCES[1]['filename'], formatting_info=True)
    groups, sheets = [], []
    for sheet in workbook.sheets():
        header_row = next((r for r in range(min(20, sheet.nrows)) if any(norm(v) == '职位代码' for v in sheet.row_values(r))), None)
        check(header_row is not None, f'No actual audit header: {sheet.name}')
        headers = sheet.row_values(header_row)
        position_column = next(c for c, v in enumerate(headers) if norm(v) == '职位代码')
        score_column = next(c for c, v in enumerate(headers) if norm(v) == '面试分数线')
        position_name_column = next(c for c, v in enumerate(headers) if norm(v) == '职位名称')
        remarks_column = next((c for c, v in enumerate(headers) if norm(v) == '备注'), None)
        unit_labels = [text(v).removeprefix('单位：') for r in range(header_row) for v in sheet.row_values(r) if text(v).startswith('单位：')]
        check(len(unit_labels) == 1, f'Ambiguous official unit label: {sheet.name}')
        unit = unit_labels[0]
        department_name = AUDIT_UNITS.get(unit)
        matches = departments.get(department_name, set()) if department_name else set()
        department_code = next(iter(matches)) if len(matches) == 1 else None
        mapping_status = 'official_unit_label_to_reviewed_unique_csv_department' if department_code else 'unresolved_department'
        # Resolve only allowed job columns within actual merged ranges, never ffill.
        safe_columns = {position_column, score_column, position_name_column}
        merges = [m for m in sheet.merged_cells if any(m[2] <= c < m[3] for c in safe_columns)]
        sheet_groups = {}
        safe_rows = 0
        for r in range(header_row + 1, sheet.nrows):
            pos_value, _ = merged_value(sheet, r, position_column, merges)
            if not text(pos_value):
                continue
            if re.fullmatch(r'\d{1,12}', text(pos_value)) is None and not isinstance(pos_value, (int, float)):
                continue
            position = code(pos_value, 12)
            cutoff_value, score_anchor = merged_value(sheet, r, score_column, merges)
            cutoff = score(cutoff_value)
            position_name, _ = merged_value(sheet, r, position_name_column, merges)
            if position not in sheet_groups:
                sheet_groups[position] = {'position_code': position, 'position_name': text(position_name), 'first_row': r+1, 'last_row': r+1, 'score_anchor_rows': set(), 'scores': set(), 'row_count': 0, 'batch_markers': set()}
            group = sheet_groups[position]
            group['row_count'] += 1
            group['last_row'] = r+1
            if cutoff is not None:
                group['scores'].add(cutoff)
                group['score_anchor_rows'].add(score_anchor+1)
            if remarks_column is not None:
                # Retain only batch tokens, never free-text remarks about a person.
                note = text(sheet.cell_value(r, remarks_column))
                group['batch_markers'].update(t for t in ['调剂', '递补'] if t in note)
            safe_rows += 1
        sheets.append({'sheet': sheet.name, 'official_unit_label': unit, 'resolved_department_name': department_name, 'resolved_department_code': department_code, 'mapping_status': mapping_status, 'header_row': header_row+1, 'cutoff_source_field': headers[score_column], 'source_job_rows': safe_rows, 'positions': len(sheet_groups)})
        for position, g in sorted(sheet_groups.items()):
            pid = key(department_code, position) if department_code else None
            groups.append({
                'id': f'audit:{sheet.name}:{position}', 'position_id': pid, 'year': YEAR,
                'batch_id': SOURCES[1]['batch_id'], 'source_id': SOURCES[1]['id'],
                'source_sheet': sheet.name, 'official_unit_label': unit,
                'department_code': department_code, 'department_name': department_name,
                'department_code_provenance': mapping_status, 'position_code': position,
                'position_name': g['position_name'], 'published_cutoff_values': [score_number(v) for v in sorted(g['scores'])],
                'source_cutoff_field': headers[score_column], 'score_unit': '分',
                'source_first_row': g['first_row'], 'source_last_row': g['last_row'],
                'source_score_anchor_rows': sorted(g['score_anchor_rows']),
                'source_list_row_count': g['row_count'], 'observed_batch_markers': sorted(g['batch_markers']),
                'cutoff_unique_in_source': len(g['scores']) == 1,
            })
    return groups, sheets


def compare_audit(audit_records, national):
    statuses = collections.Counter()
    by_position = {}
    for record in audit_records:
        pid = record['position_id']
        other = national.get(pid)
        if pid is None:
            status = 'department_unresolved'
        elif other is None:
            status = 'not_in_national_initial_list'
        elif not record['cutoff_unique_in_source'] or len(other['scores']) != 1:
            status = 'not_comparable_nonunique_or_missing_cutoff'
        else:
            official = Decimal(str(record['published_cutoff_values'][0]))
            status = 'same_published_cutoff' if official == next(iter(other['scores'])) else 'different_published_cutoff'
        record['national_initial_comparison'] = status
        record['national_initial_cutoff_values'] = [score_number(v) for v in sorted(other['scores'])] if other else None
        statuses[status] += 1
        if pid:
            check(pid not in by_position, 'Duplicate official composite position key')
            by_position[pid] = status
    return by_position, dict(sorted(statuses.items()))


def make_rows(national, positions, official_by_position):
    rows, unmatched, label_differences, conflicting, missing = [], [], [], [], []
    for pid, group in sorted(national.items()):
        raw = group['raw']
        scores = sorted(group['scores'])
        match = positions.get(pid)
        annotation = match['data'] if match else None
        if match is None:
            unmatched.append(pid)
        else:
            fields = [('招录机关', 'dept_name'), ('用人司局', 'org_name'), ('招考职位', 'position_name')]
            differences = [{ 'national_field': a, 'csv_field': b, 'national_value': raw[a], 'csv_value': annotation[b]} for a, b in fields if norm(raw[a]) != norm(annotation[b])]
            if differences:
                label_differences.append({'id': pid, 'differences': differences})
        if len(scores) > 1:
            conflicting.append({'id': pid, 'published_cutoff_values': [score_number(v) for v in scores]})
        elif not scores:
            missing.append(pid)
        attributes = []
        for _, csv_field in CSV_FIELDS:
            value = annotation[csv_field] if annotation else None
            if csv_field == 'headcount' and value is not None:
                value = int(value)
            attributes.append(value)
        rows.append([
            pid, YEAR, SOURCES[0]['batch_id'], SOURCES[0]['id'], group['source_sheet'],
            group['first_row'], group['last_row'], group['row_count'],
            raw['招录机关'], raw['部门代码'], raw['用人司局'], raw['招考职位'], raw['职位代码'],
            score_number(scores[0]) if len(scores) == 1 else None,
            [score_number(v) for v in scores], '最低面试分数', '分', sorted(group['formats']), len(scores) == 1,
            'matched_composite_key' if match else 'unmatched', match['record'] if match else None,
            *attributes, official_by_position.get(pid, 'outside_official_audit_subset'),
        ])
    return rows, unmatched, label_differences, conflicting, missing


def write_shards(directory, rows):
    directory.mkdir(parents=True, exist_ok=True)
    shards, bucket, size = [], [], 0
    def prefix(index):
        return encode({'schema_version': '1.0.0', 'dataset_id': DATASET, 'shard_index': index, 'columns': COLUMNS})[:-1] + b',"rows":['
    def flush():
        nonlocal bucket, size
        if not bucket:
            return
        index = len(shards)+1
        data = prefix(index) + b','.join(encode(row) for row in bucket) + b']}\n'
        check(len(data) < SHARD_LIMIT, 'Shard exceeds byte ceiling')
        path = directory / f'cutoffs-{index:03d}.json'
        path.write_bytes(data)
        check(json.loads(data)['rows'] == bucket, 'JSON round trip mismatch')
        shards.append({'path': path.name, 'bytes': len(data), 'sha256': digest(data), 'positions': len(bucket), 'first_id': bucket[0][0], 'last_id': bucket[-1][0]})
        bucket, size = [], 0
    for row in rows:
        row_size = len(encode(row))
        if bucket and len(prefix(len(shards)+1)) + size + row_size + len(bucket) + 3 >= SHARD_LIMIT:
            flush()
        bucket.append(row)
        size += row_size
    flush()
    expected = {s['path'] for s in shards}
    for path in directory.glob('cutoffs-*.json'):
        if path.name not in expected:
            path.unlink()  # Only obsolete files generated by this builder.
    return shards


def count_attribute(rows, name):
    index = COLUMNS.index(name)
    counts = collections.Counter(row[index] for row in rows if row[index] is not None)
    return [{'value': value, 'positions': n} for value, n in sorted(counts.items())]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workspace', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parents[1] / 'data' / 'interview_cutoffs_2026')
    args = parser.parse_args()
    verify_inputs(args.workspace)
    positions, departments = read_positions(args.workspace)
    national, national_metadata = read_national(args.workspace)
    official, official_sheets = read_audit(args.workspace, departments)
    official_by_position, comparison_counts = compare_audit(official, national)
    rows, unmatched, label_differences, conflicts, missing = make_rows(national, positions, official_by_position)
    check(all(len(row) == len(COLUMNS) for row in rows), 'Unexpected output row width')
    check(len({row[0] for row in rows}) == len(rows), 'Output composite keys are not unique')
    check(sum(row[COLUMNS.index('source_list_row_count')] for row in rows) == national_metadata['source_list_rows'], 'List-row aggregation mismatch')
    shards = write_shards(args.output, rows)
    absent = sorted(set(positions)-set(national))
    csv_source = {'source_id': SOURCES[2]['id'], 'join_key': ['year', 'dept_code_normalized_to_6_digits', 'position_code_12_digits'], 'normalization': '部门代码2000仅补前导零为002000；职位代码单独不唯一。', 'retained_annotations': dict(CSV_FIELDS), 'excluded_annotations': ['treatment_score', 'treatment_system_tier', 'treatment_region_tier', 'age_min', 'age_max', 'grassroots_years', 'min_service_years', 'extra_flags', 'remarks'], 'source_system_tag_note': '继承第三方CSV的treatment_system_class分析标签；不是官方招录系统分类。', 'source_province_note': '继承第三方CSV的work_province；原工作地点另存。'}
    statistics = {
        'schema_version': '1.0.0', 'dataset_id': DATASET, 'researched_at': RESEARCHED_AT,
        'source_list_rows': national_metadata['source_list_rows'], 'initial_list_positions': len(rows),
        'positions_with_unique_published_cutoff': len(rows)-len(conflicts)-len(missing),
        'positions_with_conflicting_published_cutoffs': len(conflicts), 'positions_without_published_cutoff': len(missing),
        'position_csv_records': len(positions), 'matched_to_position_csv': len(rows)-len(unmatched), 'unmatched_to_position_csv': len(unmatched),
        'position_csv_jobs_absent_from_initial_list': len(absent),
        'matched_positions_planned_seats': sum(row[COLUMNS.index('source_csv_planned_seats')] or 0 for row in rows),
        'official_audit_positions': len(official), 'official_audit_sheets': len(official_sheets),
        'official_comparisons': comparison_counts,
        'by_source_csv_province': count_attribute(rows, 'source_csv_province'),
        'by_source_system_tag': count_attribute(rows, 'source_system_tag'),
        'by_source_csv_education': count_attribute(rows, 'source_csv_education'),
        'by_source_csv_exam_category': count_attribute(rows, 'source_csv_exam_category'),
        'counts_are_positions_not_planned_seats_or_people': True,
        'score_aggregation_note': '未计算跨职位统一分数均值或全国分数排名；不同职位可能采用不同笔试科目与加权口径。',
    }
    validation = {
        'schema_version': '1.0.0', 'inputs_hash_verified': True,
        'source_sheet_reading': national_metadata,
        'composite_keys_unique': True, 'score_extraction': '仅源表明确的最低面试分数/面试分数线列；不读取个人考试成绩、不从人员成绩计算MIN。',
        'source_label_variant_positions': sum(bool(g['label_variants']) for g in national.values()),
        'csv_name_differences': label_differences,
        'unmatched_position_ids': unmatched, 'conflicting_published_cutoffs': conflicts, 'missing_published_cutoff_ids': missing,
        'csv_position_ids_not_listed': absent,
        'absence_note': '不在本首批名单中不等于最低线为0，也不能单凭缺失认定无人报考、无人过线或取消招录。',
        'official_audit': {'sheets': official_sheets, 'comparisons': comparison_counts, 'comparison_scope': '仅审计署这份官方公告附件覆盖岗位；差异保留，不能用后发批次覆盖全国首批值。'},
        'privacy': {'candidate_identity_columns_projected': False, 'candidate_identity_values_written': False, 'candidate_personal_scores_used': False, 'free_text_candidate_remarks_written': False, 'raw_interview_workbooks_copied_to_output': False, 'allowed_national_source_columns': ['招录机关','部门代码','用人司局','招考职位','职位代码','最低面试分数']},
        'shard_roundtrip_verified': True, 'max_shard_bytes': max(s['bytes'] for s in shards), 'all_shards_below_850000_bytes': all(s['bytes'] < 850_000 for s in shards),
    }
    files = [
        write_json(args.output/'official-audit-comparison.json', {'schema_version':'1.0.0','source_id':SOURCES[1]['id'],'records':official}),
        write_json(args.output/'statistics.json', statistics, pretty=True),
        write_json(args.output/'validation.json', validation, pretty=True),
    ]
    manifest = {
        'schema_version':'1.0.0','dataset_id':DATASET,'title':'2026国考首批逐岗位最低面试分数','year':YEAR,
        'researched_at':RESEARCHED_AT,'sources':SOURCES,'batch':{'id':SOURCES[0]['batch_id'],'label':SOURCES[0]['batch_label']},
        'record_grain':'一个招录年度、部门代码、职位代码、发布批次下的已公布最低面试分数。',
        'position_key':['year','部门代码','职位代码'],'observation_key':['year','部门代码','职位代码','batch_id'],
        'source_status':'全国首批公开镜像；审计署官方附件仅核验其覆盖的岗位，不代表全国已逐条由官方直链核验。',
        'score_definition':{'source_field':'最低面试分数','official_audit_source_field_variants':['面试分数线','面试\n分数线'],'unit':'分','unit_note':'成绩单位；保留源数值和数字格式，不推定所有岗位满分相同。','is_national_policy_passline':False,'is_actual_interview_test_score':False,'source_reported_cutoff_only':True,'cross_job_score_basis':'不同职位可能含专业科目或外语测试权重；本数据未统一换算或补造计分公式。','missing_or_conflicting_value_policy':'唯一已公布值写入最低面试分数；缺失或冲突写null，published_cutoff_values保留全部已公布不同值。'},
        'columns':COLUMNS,'column_notes':{'id':'2026:六位部门代码:十二位职位代码；本分片均属同一首批批次。','source_first_row':'Excel行号从1开始，仅作岗位证据位置。','source_list_row_count':'该岗位在源名单中的行数；没有保存人员身份，也没有据此构造个人记录。','source_cutoff_field':'源成绩字段原名。','published_cutoff_values':'源表该岗位该批次出现的不同最低线值；未选择其中最小值。','source_number_formats':'源最低线单元格的Excel数字格式。','position_csv_record':'CSV记录行号从2开始计第一条数据，仅用于复现定位。','official_audit_comparison':'审计署官方附件覆盖之外为outside_official_audit_subset，不能理解为已核验。'},
        'position_attribute_source':csv_source,'positions':len(rows),'source_list_rows':national_metadata['source_list_rows'],
        'shard_limit_bytes':SHARD_LIMIT,'shards':shards,'supporting_files':files,
        'build':{'script':'../../scripts/build_interview_cutoffs_2026.py','dependencies':['openpyxl','xlrd'],'network_required_at_build_time':False,'input_workbooks_are_temporary_and_not_repo_outputs':True},
    }
    write_json(args.output/'manifest.json',manifest,pretty=True)
    print(json.dumps({'output':str(args.output),'source_list_rows':statistics['source_list_rows'],'positions':len(rows),'matched_csv':statistics['matched_to_position_csv'],'unmatched_csv':len(unmatched),'conflicting_cutoffs':len(conflicts),'missing_cutoffs':len(missing),'official_audit_positions':len(official),'official_comparisons':comparison_counts,'shards':len(shards),'max_shard_bytes':validation['max_shard_bytes']},ensure_ascii=False))


if __name__ == '__main__':
    main()
